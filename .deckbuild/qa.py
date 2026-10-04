from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
W, H = 13.333, 7.5
prs = Presentation("C:/Users/Admin/smarttoken-dapp/SMT-Pitch-Deck.pptx")

def walk(shapes):
    for s in shapes:
        if s.shape_type == 6:
            yield from walk(s.shapes)
        else:
            yield s

issues = 0
for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in walk(slide.shapes):
        try:
            x, y = sh.left / EMU_IN, sh.top / EMU_IN
            w, h = sh.width / EMU_IN, sh.height / EMU_IN
        except TypeError:
            continue
        if x < -0.02 or y < -0.02 or x + w > W + 0.02 or y + h > H + 0.02:
            print(f"[S{idx}] OUT OF CANVAS: ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f} '{(sh.text_frame.text[:40] if sh.has_text_frame else sh.shape_type)}'")
            issues += 1
        if sh.has_text_frame and sh.text_frame.text.strip():
            # crude overflow estimate: chars per line vs box width & height
            tf = sh.text_frame
            lines = 0
            for p in tf.paragraphs:
                txt = p.text
                sizes = [r.font.size.pt for r in p.runs if r.font.size]
                fs = max(sizes) if sizes else 18
                # avg glyph width ~ 0.55 * fontsize for Latin (Segoe UI)
                cpl = max(int(w * 72 / (0.55 * fs)), 1)
                import math
                n = max(math.ceil(len(txt) / cpl), 1) if txt else 1
                lines += n
                tot_h = 0
            # recompute height
            total_h = 0
            for p in tf.paragraphs:
                txt = p.text
                sizes = [r.font.size.pt for r in p.runs if r.font.size]
                fs = max(sizes) if sizes else 18
                cpl = max(int(w * 72 / (0.55 * fs)), 1)
                import math
                n = max(math.ceil(len(txt) / cpl), 1) if txt else 1
                total_h += n * fs * 1.35 / 72
            if total_h > h * 1.25 and h > 0.2:
                print(f"[S{idx}] POSSIBLE OVERFLOW: box {w:.2f}x{h:.2f} est_text_h={total_h:.2f} '{tf.text[:60].replace(chr(10),' | ')}'")
                issues += 1
            boxes.append((x, y, w, h, tf.text[:30]))
print(f"done, {issues} issues, {len(prs.slides)} slides")
